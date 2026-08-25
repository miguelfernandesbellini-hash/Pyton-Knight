import json
from collections import defaultdict
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'Documentacao_Evolucao_Pyton_Knight_COMPLETA_ATUALIZADA.docx'
ACTIVITIES = json.loads((ROOT / 'build/activity_catalog.json').read_text(encoding='utf-8'))
RESULTS = json.loads((ROOT / 'build/test_results.json').read_text(encoding='utf-8'))

# Preset único: compact_reference_guide. Override nomeado: capa editorial.
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
BLACK = RGBColor(0, 0, 0)
MUTED = RGBColor(0x66, 0x6B, 0x73)
WIDTH = 9360
INDENT = 120
FILL = 'E8EEF5'

MECHANICS = {
    'M01':'Placa de Pressão','M02':'Placa Rúnica Alternadora','M03':'Alavanca',
    'M05':'Porta Simples','M06':'Porta Condicional','M07':'Ponte Retrátil',
    'M08':'Ponte Segmentada','M09':'Espinhos','M12':'Armadilha Cíclica',
    'M14':'Espelhos Coloridos','M17':'Pedestal de Entrada','M18':'Runa de Saída',
    'M19':'Cofre Rúnico','M20':'Guardião / Validador','M21':'Chave','M22':'Baú',
    'M23':'Sequência de Baús / Busca','M24':'Moeda / Rubi','M26':'Totem / Runa Sequencial',
    'M30':'Sensores do Mundo','M31':'Movimento Orientado','M32':'Execução e Reset',
    'M33':'Sistema de Vidas','M34':'Bolsa Persistente','M35':'XP por Desempenho',
    'M36':'Objetivos e Restrições','M37':'Orçamento de Código'
}

def font(run, size=11, color=BLACK, bold=False, italic=False, name='Calibri'):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn('w:ascii'), name); fonts.set(qn('w:hAnsi'), name); fonts.set(qn('w:eastAsia'), name)
    run.font.size = Pt(size); run.font.color.rgb = color; run.bold = bold; run.italic = italic

def shade(cell, fill=FILL):
    tcpr = cell._tc.get_or_add_tcPr(); node = tcpr.find(qn('w:shd'))
    if node is None: node = OxmlElement('w:shd'); tcpr.append(node)
    node.set(qn('w:fill'), fill); node.set(qn('w:val'), 'clear')

def margins(cell):
    tcpr = cell._tc.get_or_add_tcPr(); mar = tcpr.find(qn('w:tcMar'))
    if mar is None: mar = OxmlElement('w:tcMar'); tcpr.append(mar)
    for edge, value in [('top',80),('bottom',80),('start',120),('end',120)]:
        node = mar.find(qn(f'w:{edge}'))
        if node is None: node = OxmlElement(f'w:{edge}'); mar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def table_geometry(table, widths):
    assert sum(widths) == WIDTH
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.LEFT
    pr = table._tbl.tblPr
    for tag, value in [('tblW', WIDTH), ('tblInd', INDENT)]:
        node = pr.find(qn(f'w:{tag}'))
        if node is None: node = OxmlElement(f'w:{tag}'); pr.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')
    layout = pr.find(qn('w:tblLayout'))
    if layout is None: layout = OxmlElement('w:tblLayout'); pr.append(layout)
    layout.set(qn('w:type'), 'fixed')
    borders = OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        b = OxmlElement(f'w:{edge}'); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'B7C2CE'); borders.append(b)
    pr.append(borders)
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for value in widths:
        col = OxmlElement('w:gridCol'); col.set(qn('w:w'), str(value)); grid.append(col)
    for rindex, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if rindex == 0:
            repeat = OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); trpr.append(repeat)
        for cell, value in zip(row.cells, widths):
            margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr(); tcw = tcpr.find(qn('w:tcW'))
            if tcw is None: tcw = OxmlElement('w:tcW'); tcpr.append(tcw)
            tcw.set(qn('w:w'), str(value)); tcw.set(qn('w:type'),'dxa')
            if rindex == 0: shade(cell)

def add_table(doc, headers, rows, widths, trailing=True):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, value in enumerate(headers):
        p = table.rows[0].cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(2)
        font(p.add_run(str(value)), 8.5, DARK_BLUE, True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]; p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
            font(p.add_run(str(value)), 8.5)
    table_geometry(table, widths)
    if trailing:
        doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_numbering(doc, kind):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn('w:abstractNumId'))) for x in numbering.findall(qn('w:abstractNum'))]
    aid = max(abstract_ids, default=-1) + 1
    abstract = OxmlElement('w:abstractNum'); abstract.set(qn('w:abstractNumId'), str(aid))
    multi = OxmlElement('w:multiLevelType'); multi.set(qn('w:val'), 'multilevel' if kind == 'heading' else 'singleLevel'); abstract.append(multi)
    levels = 3 if kind == 'heading' else 1
    for level in range(levels):
        lvl = OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'), str(level))
        start = OxmlElement('w:start'); start.set(qn('w:val'),'1'); lvl.append(start)
        fmt = OxmlElement('w:numFmt'); fmt.set(qn('w:val'), 'decimal' if kind != 'bullet' else 'bullet'); lvl.append(fmt)
        txt = OxmlElement('w:lvlText')
        txt.set(qn('w:val'), ('•' if kind == 'bullet' else ('.'.join(f'%{x}' for x in range(1,level+2)) + '.'))); lvl.append(txt)
        if kind == 'heading':
            ps = OxmlElement('w:pStyle'); ps.set(qn('w:val'), f'Heading{level+1}'); lvl.append(ps)
        ppr = OxmlElement('w:pPr'); tabs = OxmlElement('w:tabs'); tab = OxmlElement('w:tab'); tab.set(qn('w:val'),'num'); tab.set(qn('w:pos'),'720' if kind == 'heading' else '540'); tabs.append(tab); ppr.append(tabs)
        ind = OxmlElement('w:ind'); ind.set(qn('w:left'),'720' if kind == 'heading' else '540'); ind.set(qn('w:hanging'),'720' if kind == 'heading' else '271'); ppr.append(ind); lvl.append(ppr); abstract.append(lvl)
    numbering.append(abstract)
    num_ids = [int(x.get(qn('w:numId'))) for x in numbering.findall(qn('w:num'))]; nid = max(num_ids, default=0) + 1
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(nid)); ref = OxmlElement('w:abstractNumId'); ref.set(qn('w:val'), str(aid)); num.append(ref); numbering.append(num)
    return nid

def number_paragraph(p, nid, level=0):
    ppr = p._p.get_or_add_pPr(); numpr = OxmlElement('w:numPr'); ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'),str(level)); numid = OxmlElement('w:numId'); numid.set(qn('w:val'),str(nid)); numpr.append(ilvl); numpr.append(numid); ppr.append(numpr)

def setup(doc):
    sec = doc.sections[0]; sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.top_margin = sec.right_margin = sec.bottom_margin = sec.left_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492); sec.different_first_page_header_footer = True
    normal = doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11); normal.font.color.rgb=BLACK
    normal._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
    normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.25
    for name, size, color, before, after, page in [('Heading 1',16,BLUE,18,10,True),('Heading 2',13,BLUE,14,7,False),('Heading 3',12,DARK_BLUE,10,5,False)]:
        s=doc.styles[name]; s.font.name='Calibri'; s.font.size=Pt(size); s.font.color.rgb=color; s.font.bold=True
        s._element.rPr.rFonts.set(qn('w:ascii'),'Calibri'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri')
        s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after); s.paragraph_format.keep_with_next=True; s.paragraph_format.page_break_before=page
    h=sec.header.paragraphs[0]; h.paragraph_format.tab_stops.add_tab_stop(Inches(6.5),WD_TAB_ALIGNMENT.RIGHT)
    font(h.add_run('PYTON KNIGHT · DOCUMENTAÇÃO DE EVOLUÇÃO'),8.5,MUTED,True); font(h.add_run('\t20 ATIVIDADES'),8.5,MUTED)
    sec.first_page_header.paragraphs[0].text=''
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT; font(f.add_run('Página '),8.5,MUTED)
    run=f.add_run(); begin=OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin'); instr=OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text='PAGE'; end=OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end'); run._r.extend([begin,instr,end]); font(run,8.5,MUTED)
    ff=sec.first_page_footer.paragraphs[0]; ff.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(ff.add_run('Edição consolidada · 22 de agosto de 2026'),8.5,MUTED)

def heading(doc, text, level, nid):
    p=doc.add_paragraph(style=f'Heading {level}'); font(p.add_run(text), {1:16,2:13,3:12}[level], BLUE if level<3 else DARK_BLUE, True); number_paragraph(p,nid,level-1); return p

def body(doc, text):
    p=doc.add_paragraph(); font(p.add_run(text)); return p

def bullet(doc, text, nid):
    p=doc.add_paragraph(); number_paragraph(p,nid); p.paragraph_format.space_after=Pt(4); p.paragraph_format.line_spacing=1.25; font(p.add_run(text)); return p

def cover(doc):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(58)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(14); font(p.add_run('DOCUMENTAÇÃO TÉCNICA E PEDAGÓGICA'),10,BLUE,True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(9); font(p.add_run('Evolução do Pyton Knight'),27,DARK_BLUE,True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after=Pt(20); font(p.add_run('Da Unidade 1 à edição completa de 20 atividades'),15,BLUE)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.left_indent=Inches(.55); p.paragraph_format.right_indent=Inches(.55); p.paragraph_format.space_after=Pt(80)
    font(p.add_run('Histórico preservado, arquitetura modular, Livro Mágico, Python controlado, mecânicas M01–M37 utilizadas e evidências de teste.'),11)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run('VERSÃO COMPLETA E ATUALIZADA'),11,DARK_BLUE,True)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; font(p.add_run('22 de agosto de 2026'),10,MUTED)

def build():
    doc=Document(); setup(doc); hn=add_numbering(doc,'heading'); bn=add_numbering(doc,'bullet'); cover(doc)
    heading(doc,'Resumo executivo',1,hn)
    body(doc,'Esta edição consolida o Pyton Knight como jogo educacional de quatro unidades e 20 atividades. A implementação preserva a base Phaser e a Unidade 1, migra o movimento para orientação relativa e acrescenta interpretador seguro, Livro Mágico permanente, objetivos declarativos, vidas, orçamento, XP, moedas e persistência.')
    body(doc,'O aceite lógico registrou 25 cenários oficiais/variantes, 300 verificações obrigatórias no total (15 por atividade), 25 reexecuções após reset e 19 regressões globais, todos com PASS. A homologação visual em navegador real permanece registrada como pendência real.')
    add_table(doc,['Indicador','Resultado','Evidência'],[
        ('Unidades','4','Progressão 1–5, 6–10, 11–15 e 16–20'),('Atividades','20','Catálogo declarativo em activities.js'),('Cenários oficiais','25 PASS','Inclui variantes de 12, 19 e 20'),('Matriz obrigatória','300 PASS','15 verificações em cada atividade'),('Reexecuções','25 PASS','Código preservado e recompensa estável'),('Regressões','19 PASS','Parser, runtime, vidas e persistência'),('Mecânicas','27 códigos','Necessárias + M31–M37')],[2300,1700,5360])

    heading(doc,'Histórico preservado da Unidade 1',1,hn)
    heading(doc,'Baseline recebido',2,hn)
    body(doc,'O ZIP de origem continha 19 arquivos na raiz, Phaser local, cinco assets, cenas de inicialização e módulos de barreira, movimento, interpretação, renderização e interface. O catálogo possuía cinco atividades. Essa base foi evoluída, não recriada.')
    for text in ['Phaser 3, cenas, assets originais e Preloader foram preservados.','Mapas, barreiras, cristais e sequência pedagógica das Atividades 1–5 foram mantidos.','A estrutura systems/ e ui/ existente foi ampliada.','A execução continua sem build e sem dependências de produção.']: bullet(doc,text,bn)
    heading(doc,'Migração necessária',2,hn)
    add_table(doc,['Aspecto','Base recebida','Edição atual'],[
        ('Movimento','direções absolutas/acopladas ao mapa','avanço relativo à orientação + giros'),('Execução','reset parcial','tentativa reconstruída antes da AST'),('Livro Mágico','interface da Unidade 1','painel permanente com tutor, I/O e objetivos'),('Interpretador','variáveis e aritmética','subconjunto progressivo até loops'),('Conclusão','chegada como critério dominante','objetivos físicos e de código'),('Falhas','término pouco específico','causas independentes')],[1800,3300,4260])

    heading(doc,'Fontes oficiais e regra de conflito',1,hn)
    for text in ['Prompt Mestre Work atualizado.','Biblioteca Oficial de Mecânicas v2.','Especificação das Atividades 6 a 20 v2.','Especificação do Livro Mágico.','Roteiro Pedagógico Definitivo v3.','Código do ZIP GitHub-Ready como baseline.']: bullet(doc,text,bn)
    body(doc,'Em conflitos, foram priorizados: preservar o que funciona; escolher a solução mais simples, modular e testável; respeitar a progressão pedagógica; evitar dependências; implementar somente mecânicas necessárias.')

    heading(doc,'Arquitetura final',1,hn)
    add_table(doc,['Módulo','Responsabilidade'],[
        ('activities.js','20 atividades, mapas, objetivos, tutoriais, limites e soluções.'),('PythonSubsetParser.js','Tokenização, indentação, AST e análise semântica.'),('CommandInterpreter.js','Validação pedagógica e execução assíncrona.'),('DungeonSystem.js','API, sensores, entidades e conexões.'),('PlayerController.js','Orientação, movimento, colisão e perigos.'),('MissionObjectiveSystem.js','Objetivos físicos e computacionais.'),('TutorialSystem.js','Microetapas introdutórias.'),('CodeBudgetValidator.js','Contagem semântica.'),('PersistenceService.js','Estado persistente versionado.'),('ProgressionSystem.js','Desbloqueio e recompensas.'),('MapRenderer.js / BarrierSystem.js','Mapa, barreiras e estados visuais.'),('GameUI.js','Livro Mágico e HUD.')],[3000,6360])
    heading(doc,'Fluxo de EXECUTAR',2,hn)
    for text in ['Preservar o texto do editor.','Restaurar posição, orientação, entidades e itens transitórios.','Gerar AST e análise semântica.','Validar conceito, whitelist e orçamento.','Executar com guardas de loop.','Avaliar objetivos e causa de término.','Consolidar progresso apenas após vitória válida.']: bullet(doc,text,bn)

    heading(doc,'Livro Mágico',1,hn)
    body(doc,'O Livro Mágico permanece aberto durante a atividade e não foi transformado em menu ou enciclopédia separada.')
    add_table(doc,['Área','Comportamento'],[
        ('Missão','nome, descrição e conceito atual'),('Objetivos','estado de cada condição'),('Tutor','microetapas em 1, 6, 11 e 16'),('Editor','sintaxe Python e código persistente entre execuções'),('Entrada','campo ativado quando input() pausa'),('Console','saídas de print()'),('Feedback','mensagem ligada à causa real'),('HUD','vidas, orientação, XP, moedas e progresso'),('Ações','EXECUTAR, RESTAURAR e navegação')],[1900,7460])

    heading(doc,'Movimento orientado e colisões',1,hn)
    body(doc,'Guto possui orientação NORTE, LESTE, SUL ou OESTE. andar_frente() usa o vetor atual; os giros mudam apenas a orientação. Cada EXECUTAR restaura a orientação inicial declarada.')
    add_table(doc,['Comando','Efeito','Validação'],[
        ('andar_frente(passos=1)','avança na direção atual','limite, parede, grade, porta e perigo por casa'),('virar_direita()','gira 90° horário','não desloca'),('virar_esquerda()','gira 90° anti-horário','não desloca')],[2450,2600,4310])
    heading(doc,'Falsa colisão corrigida',2,hn)
    body(doc,'O fim de um programa que não cumpriu a missão retorna INCOMPLETE_EXECUTION. WALL_COLLISION ocorre somente em tentativa física de atravessar parede/limite. Porta fechada e morte letal possuem causas próprias.')

    heading(doc,'Interpretador Python controlado',1,hn)
    body(doc,'O parser reconhece somente a gramática necessária, produz AST e registra conceitos, comandos e instruções semânticas. Não há eval(), new Function() nem execução de Python arbitrário.')
    add_table(doc,['Unidade','Recursos'],[('1','variáveis, textos, booleanos, aritmética e movimento'),('2','print(), input(), int() e processamento'),('3','if, elif, else, comparadores, and, or, not e sensores'),('4','for, range(), acumuladores, while e break')],[1500,7860])
    for text in ['Whitelist para API da dungeon.','Bloqueio de conceitos futuros.','Validação de tipos e divisão por zero.','Limite de 600 instruções e 100 iterações.','Erros vinculados à linha quando disponível.','Soluções alternativas validadas por estado e objetivos.']: bullet(doc,text,bn)

    heading(doc,'Execução, reset e estados',1,hn)
    add_table(doc,['Estado da tentativa','Estado persistente'],[
        ('posição, orientação, entidades, portas, itens, console e variáveis','conclusões, desbloqueio, XP, bolsa e recompensas'),('recriado antes de EXECUTAR','mantido entre execuções e sessões'),('não inclui o texto do editor','não guarda tentativa parcial')],[4680,4680])
    for text in ['Código do aluno permanece.','Vidas não reiniciam em cada EXECUTAR.','Mecanismos temporários retornam ao estado inicial.','XP, moedas e desbloqueios não são apagados.']: bullet(doc,text,bn)

    heading(doc,'Objetivos, restrições e orçamento',1,hn)
    body(doc,'Os objetivos combinam chegada, saída, conceitos, comandos, entidades, flags, chave, moedas, variáveis e orçamento. Progresso parcial é mostrado sem ser confundido com colisão.')
    budgets=[(a['id'],a['nome'],a['instructionBudget']) for a in ACTIVITIES if a.get('instructionBudget')]
    add_table(doc,['Atividade','Nome','Limite'],budgets,[1200,6660,1500])
    body(doc,'Comentários e linhas vazias não contam; ponto e vírgula gera instruções reais.')

    heading(doc,'Vidas, moedas, XP e persistência',1,hn)
    add_table(doc,['Sistema','Regra'],[
        ('Vidas','3 por atividade; apenas HAZARD_DEATH remove vida.'),('Moedas/rubis','coleta provisória, crédito após vitória válida.'),('XP','base da atividade × 1,00 / 0,75 / 0,50 por vidas.'),('Replay','executável sem duplicar recompensa.'),('Persistência','localStorage, esquema v2 e normalização defensiva.')],[2100,7260])
    body(doc,'rewardedActivities registra a consolidação por atividade e completionProcessed impede duplicação na mesma execução. As 25 reexecuções confirmaram estabilidade.')

    heading(doc,'Progressão pedagógica e atividades',1,hn)
    add_table(doc,['Unidade','Atividades','Conteúdo','Tutorial','Boss'],[('1','1–5','variáveis e orientação','1','5'),('2','6–10','entrada, saída e processamento','6','10'),('3','11–15','condicionais e sensores','11','15'),('4','16–20','repetições e busca','16','20')],[1200,1500,3860,1400,1400])
    for unit in range(1,5):
        heading(doc,f'Unidade {unit}',2,hn); rows=[]
        for a in [x for x in ACTIVITIES if x['unidade']==unit]:
            profile=[]
            if a.get('tutorial'): profile.append('Tutorial')
            if a['id']%5==0: profile.append('Boss')
            if a.get('instructionBudget'): profile.append(f"Limite {a['instructionBudget']}")
            rows.append((a['id'],a['nome'],a['concept'],', '.join(profile) or 'Aplicação'))
        add_table(doc,['#','Atividade','Conceito','Perfil'],rows,[700,2600,4260,1800])
    body(doc,'Atividades introdutórias usam microetapas; as demais não entregam a solução completa. Bosses mantêm múltiplos subproblemas e objetivos.')

    heading(doc,'Mecânicas implementadas e reutilizadas',1,hn)
    use=defaultdict(list)
    for a in ACTIVITIES:
        for code in a['mechanics']: use[code].append(a['id'])
    rows=[]
    for code in sorted(use,key=lambda x:int(x[1:])):
        ids='1–20' if len(use[code])==20 else ', '.join(map(str,use[code])); rows.append((code,MECHANICS.get(code,'Mecânica configurável'),ids))
    add_table(doc,['Código','Mecânica / sistema','Atividades'],rows,[1200,5560,2600],trailing=False)

    heading(doc,'Estratégia e resultados de testes',1,hn)
    body(doc,'A suíte executa parser e runtime em cenas de teste. Cada atividade percorre 15 verificações obrigatórias; uma solução só passa quando cumpre os objetivos e retorna SUCCESS.')
    by=defaultdict(list)
    for item in RESULTS['officialSolutions']: by[item['activity']].append(item)
    rows=[]
    matrix={x['activity']:x for x in RESULTS['activityMatrix']}
    for a in ACTIVITIES:
        cases=by[a['id']]; repeat=all(x.get('replayCause')=='SUCCESS' and x.get('codePreserved') and x.get('rewardStableOnReplay') for x in cases)
        rows.append((a['id'],len(cases),'PASS' if all(x['cause']=='SUCCESS' for x in cases) else 'FALHA',matrix[a['id']]['status'],'PASS' if repeat else 'FALHA',a.get('instructionBudget','N/A')))
    add_table(doc,['Ativ.','Casos','Solução','Matriz 15','Replay','Limite'],rows,[900,1000,1500,1500,2000,2460])
    heading(doc,'Regressões globais',2,hn)
    labels={'incompleteExecution':'Código incompleto não vira colisão','wallCollision':'Colisão real','closedDoorBlock':'Porta fechada','hazardLifeLoss':'Perigo remove uma vida','codePreservedAfterDeath':'Código após morte','unknownCommand':'Comando desconhecido','lockedConcept':'Conceito futuro bloqueado','loopGuard':'Loop infinito interrompido','syntaxError':'Erro de sintaxe','semanticTypeError':'Erro de tipos','semanticBudgetCount':'Ponto e vírgula contado','indexAssets':'Scripts e cinco assets presentes','noAbsoluteMovementApi':'Sem movimento absoluto','noDynamicCodeExecution':'Sem eval/new Function','budgetExceededNoLifeLoss':'Orçamento sem vida','deterministicRerun':'Reexecução determinística','guidedTutorials':'Tutoriais completos','rewardIdempotency':'Recompensas estáveis','xpMonotonicity':'XP por vidas'}
    add_table(doc,['Verificação','Resultado'],[(labels.get(k,k),v) for k,v in RESULTS['regressionChecks'].items()],[7860,1500])
    body(doc,'As 300 verificações da matriz funcional (15 por atividade), as regressões, a sintaxe e o carregamento HTTP foram validados. Como não havia binário de navegador no ambiente, a passagem visual das 20 telas permanece PENDENTE e não foi convertida artificialmente em PASS.')

    heading(doc,'Problemas, causas e soluções',1,hn)
    add_table(doc,['Problema','Causa','Solução'],[
        ('Falsa colisão','fim do código e falha física misturados','causas independentes'),('Estado residual','reset parcial','resetRun reconstrói a tentativa'),('Movimento absoluto','API acoplada ao mapa','orientação e comandos relativos'),('Recompensa duplicada','sem registro por atividade','rewardedActivities'),('Loop infinito','laço sem guarda','limites de execução'),('Alternativas rejeitadas','comparação textual','AST + objetivos por estado'),('Lógica duplicada','mecânicas ad hoc','entidades declarativas'),('Aritmética inválida','coerção implícita','validação de tipos')],[2600,3100,3660])

    heading(doc,'Balanceamento provisório',1,hn)
    add_table(doc,['Parâmetro','Valor atual','Local'],[
        ('XP-base padrão','100; bases por atividade','GameConstants / activities'),('XP-base boss','160','GameConstants'),('Multiplicadores','3: 1,00 · 2: 0,75 · 1: 0,50','GameConstants'),('Moeda','1','GameConstants'),('Replay','sem recompensa','GameConstants'),('Orçamentos','9 a 25','activities'),('Guardas','600 instruções; 100 iterações','GameConstants')],[2900,3300,3160])

    heading(doc,'Limitações e pendências reais',1,hn)
    for text in ['Subconjunto educacional, não Python completo.','Persistência local ao navegador.','Símbolos provisórios para algumas mecânicas.','Áudio e animações avançadas fora do escopo.','Homologação visual manual necessária.','Rebalanceamento após playtests.']: bullet(doc,text,bn)

    heading(doc,'Mapa de arquivos',1,hn)
    add_table(doc,['Grupo','Arquivos'],[('Núcleo','Game.js, MainMenu.js, index.html, activities.js'),('Interpretação','PythonSubsetParser, CommandInterpreter, CodeBudgetValidator'),('Mundo','DungeonSystem, PlayerController, BarrierSystem, MapRenderer'),('Missão','MissionObjectiveSystem, TutorialSystem'),('Progresso','PersistenceService, ProgressionSystem, GameConstants'),('Interface','ui/GameUI.js'),('Testes','tests/run-tests.cjs'),('Documentos','README, gabarito e relatórios')],[2200,7160])

    heading(doc,'Conclusão',1,hn)
    body(doc,'A evolução mantém a identidade e a base funcional do Pyton Knight, aplica as decisões oficiais e amplia o projeto para 20 atividades coerentes. O runtime modular valida resultado e estado, protege o navegador contra loops e diferencia falhas pedagógicas de perigos letais.')
    body(doc,'Os 25 cenários oficiais, as 300 verificações da matriz funcional (15 por atividade) e as regressões estão aprovados. A pendência de homologação é visual, acompanhada de futuros playtests para arte e balanceamento. Código, gabarito, README e relatórios permanecem alinhados ao mesmo catálogo e às mesmas evidências.')
    doc.core_properties.title='Documentação de Evolução do Pyton Knight — Completa e Atualizada'; doc.core_properties.author='Equipe Pyton Knight'; doc.core_properties.subject='Arquitetura, pedagogia, atividades e testes'
    doc.save(OUT); return doc

def audit(doc):
    sec=doc.sections[0]
    assert int(sec.page_width)==int(Inches(8.5)) and int(sec.page_height)==int(Inches(11))
    assert all(int(x)==int(Inches(1)) for x in [sec.top_margin,sec.right_margin,sec.bottom_margin,sec.left_margin])
    assert len(ACTIVITIES)==20 and OUT.exists() and OUT.stat().st_size>20000
    for t in doc.tables:
        assert t._tbl.tblPr.find(qn('w:tblW')).get(qn('w:w'))==str(WIDTH)
        assert t._tbl.tblPr.find(qn('w:tblInd')).get(qn('w:w'))==str(INDENT)
    text='\n'.join(p.text for p in doc.paragraphs); assert '====' not in text and 'PLACEHOLDER' not in text

if __name__=='__main__':
    document=build(); audit(document)
    print(json.dumps({'status':'PASS','output':str(OUT),'tables':len(document.tables),'paragraphs':len(document.paragraphs)},ensure_ascii=False,indent=2))
